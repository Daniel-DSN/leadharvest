"""Export leads and rejected rows to CSV, JSON or styled Excel."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from .models import Lead, Rejected

FIELDS = ["name", "email", "phone", "website", "company", "city", "domain"]
REJECTED_FIELDS = ["name", "email", "reason"]
SUPPORTED_SUFFIXES = {".csv", ".json", ".xlsx", ".xlsm"}


def to_csv(leads: list[Lead], path: Path) -> Path:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=FIELDS)
        writer.writeheader()
        for lead in leads:
            writer.writerow(lead.to_dict())
    return path


def to_json(leads: list[Lead], path: Path) -> Path:
    payload = [lead.to_dict() for lead in leads]
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def to_xlsx(leads: list[Lead], path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "leads"
    ws.append(FIELDS)
    for lead in leads:
        ws.append([lead.to_dict()[f] for f in FIELDS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def export(leads: list[Lead], path: str | Path) -> Path:
    """Write leads to ``.csv`` / ``.json`` / ``.xlsx`` (chosen by suffix)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return to_csv(leads, path)
    if suffix == ".json":
        return to_json(leads, path)
    if suffix in {".xlsx", ".xlsm"}:
        return to_xlsx(leads, path)
    raise ValueError(
        f"unsupported format: {suffix!r} (use .csv, .json or .xlsx)"
    )


def write_rejected(rejected: list[Rejected], path: str | Path) -> Path:
    """Write the rejection report (CSV by default, suffix-dispatched)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".json":
        payload = [row.to_dict() for row in rejected]
        path.write_text(
            json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        return path
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=REJECTED_FIELDS)
        writer.writeheader()
        for row in rejected:
            writer.writerow(row.to_dict())
    return path


def from_csv(path: Path) -> list[Lead]:
    with path.open(newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        return [Lead.from_dict(dict(row)) for row in reader]


def from_json(path: Path) -> list[Lead]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, list):
        raise ValueError(f"{path}: expected a JSON list of leads")
    return [Lead.from_dict(row) for row in data]


def from_xlsx(path: Path) -> list[Lead]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return []
    keys = [str(cell) if cell is not None else "" for cell in header]
    return [
        Lead.from_dict(
            {k: ("" if v is None else v) for k, v in zip(keys, row) if k}
        )
        for row in rows
    ]


def load_leads(path: str | Path) -> list[Lead]:
    """Read back an exported lead file (``.csv`` / ``.json`` / ``.xlsx``)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return from_csv(path)
    if suffix == ".json":
        return from_json(path)
    if suffix in {".xlsx", ".xlsm"}:
        return from_xlsx(path)
    raise ValueError(
        f"unsupported format: {suffix!r} (use .csv, .json or .xlsx)"
    )
