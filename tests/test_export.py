"""Export and CLI end-to-end tests — offline, using tmp_path."""

from pathlib import Path

import pytest

from leadharvest.cli import main
from leadharvest.enrich import enrich_records
from leadharvest.export import export, load_leads, write_rejected
from leadharvest.models import Lead
from leadharvest.sources import load_records

ROOT = Path(__file__).parent.parent
CSV_SAMPLE = ROOT / "samples" / "leads_raw.csv"


def make_leads() -> list[Lead]:
    return enrich_records(load_records(CSV_SAMPLE)).leads


def test_export_csv_roundtrip(tmp_path):
    leads = make_leads()
    out = export(leads, tmp_path / "leads.csv")
    header = out.read_text(encoding="utf-8").splitlines()[0]
    assert header == "name,email,phone,website,company,city,domain"
    reloaded = load_leads(out)
    assert [lead.email for lead in reloaded] == [lead.email for lead in leads]


def test_export_json_roundtrip(tmp_path):
    leads = make_leads()
    out = export(leads, tmp_path / "leads.json")
    reloaded = load_leads(out)
    assert reloaded[0].to_dict() == leads[0].to_dict()
    assert len(reloaded) == len(leads)


def test_export_xlsx_roundtrip_and_formatting(tmp_path):
    from openpyxl import load_workbook

    leads = make_leads()
    out = export(leads, tmp_path / "leads.xlsx")
    wb = load_workbook(out)
    ws = wb.active
    assert ws.title == "leads"
    assert ws.freeze_panes == "A2"
    assert ws.cell(row=1, column=1).font.bold is True
    assert ws.max_row == len(leads) + 1
    reloaded = load_leads(out)
    assert [lead.email for lead in reloaded] == [lead.email for lead in leads]


def test_export_rejects_unknown_suffix(tmp_path):
    with pytest.raises(ValueError, match="unsupported format"):
        export([], tmp_path / "leads.pdf")


def test_load_leads_rejects_unknown_suffix(tmp_path):
    path = tmp_path / "leads.txt"
    path.write_text("name\nAna\n", encoding="utf-8")
    with pytest.raises(ValueError, match="unsupported format"):
        load_leads(path)


def test_write_rejected_report(tmp_path):
    result = enrich_records(load_records(CSV_SAMPLE))
    out = write_rejected(result.rejected, tmp_path / "rejected.csv")
    lines = out.read_text(encoding="utf-8").splitlines()
    assert lines[0] == "name,email,reason"
    assert len(lines) == 1 + result.rejected_count
    assert any("invalid email format" in line for line in lines)
    assert any("missing email" in line for line in lines)


def test_enrich_cli_happy_path(tmp_path):
    out = tmp_path / "leads.xlsx"
    report = tmp_path / "rejected.csv"
    code = main(
        ["enrich", str(CSV_SAMPLE), "-o", str(out), "--report", str(report)]
    )
    assert code == 0
    assert out.exists() and report.exists()
    assert len(load_leads(out)) == 9


def test_enrich_cli_zero_valid_leads_exits_2(tmp_path):
    raw = tmp_path / "raw.csv"
    raw.write_text(
        "name,email,phone,website,company,city\nBroken,not-an-email,,,,\n",
        encoding="utf-8",
    )
    out = tmp_path / "leads.csv"
    code = main(["enrich", str(raw), "-o", str(out)])
    assert code == 2
    assert not out.exists()


def test_stats_cli_prints_totals(tmp_path, capsys):
    xlsx = export(make_leads(), tmp_path / "leads.xlsx")
    code = main(["stats", str(xlsx)])
    captured = capsys.readouterr()
    assert code == 0
    assert "9 valid leads" in captured.out
    assert "by city:" in captured.out
    assert "by domain (top 5):" in captured.out
    assert "agencia.com.br" in captured.out


def test_stats_cli_empty_file_exits_2(tmp_path, capsys):
    out = tmp_path / "leads.csv"
    export([], out)
    code = main(["stats", str(out)])
    assert code == 2
    assert "zero valid leads" in capsys.readouterr().err
